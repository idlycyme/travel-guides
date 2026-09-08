import sys
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, out = sys.argv[1], sys.argv[2]
soup = BeautifulSoup(open(src, encoding='utf-8').read(), 'html.parser')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.5)

style = doc.styles['Normal']
style.font.name = 'Arial'
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang TC')
style.font.size = Pt(10.5)

def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '1155CC'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    new_run.append(rPr); t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    new_run.append(t); hl.append(new_run); paragraph._p.append(hl)

def render_inline(par, node, bold=False, color=None, size=None):
    """Render inline children of node into paragraph par."""
    kids=[c for c in node.children if not (isinstance(c, NavigableString) and str(c).strip()=='' and not isinstance(c, Tag))] if False else list(node.children)
    # drop whitespace-only text at the very start/end
    while kids and isinstance(kids[0], NavigableString) and str(kids[0]).strip()=='': kids.pop(0)
    while kids and isinstance(kids[-1], NavigableString) and str(kids[-1]).strip()=='': kids.pop()
    for child in kids:
        if isinstance(child, NavigableString):
            txt = str(child).replace('\n', ' ')
            if txt.strip() == '' and txt == '': continue
            if txt.strip() == '': txt = ' '
            run = par.add_run(txt); run.bold = bold
            if color: run.font.color.rgb = RGBColor.from_string(color)
            if size: run.font.size = Pt(size)
        elif isinstance(child, Tag):
            if child.name == 'br':
                par.add_run().add_break(WD_BREAK.LINE)
            elif child.name == 'a':
                add_hyperlink(par, child.get('href', ''), child.get_text())
            elif child.name == 'b':
                s = child.get('style', '')
                sz = 13 if '13pt' in s else size
                render_inline(par, child, bold=True, color=color, size=sz)
            elif child.name == 'span':
                cls = child.get('class', [])
                col = 'BB0000' if 'todo' in cls else ('555555' if 'note' in cls else color)
                render_inline(par, child, bold=bold, color=col, size=size)
            else:
                render_inline(par, child, bold=bold, color=color, size=size)

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)

body = soup.body
for el in body.children:
    if not isinstance(el, Tag): continue
    if el.name == 'h1':
        p = doc.add_paragraph(); r = p.add_run(el.get_text(strip=True)); r.bold = True; r.font.size = Pt(16)
    elif el.name == 'h2':
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12);         r = p.add_run(el.get_text(strip=True)); r.bold = True; r.font.size = Pt(13)
    elif el.name == 'p':
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        render_inline(p, el)
    elif el.name == 'table':
        rows = el.find_all('tr')
        ncols = len(rows[0].find_all(['th', 'td']))
        t = doc.add_table(rows=0, cols=ncols); t.style = 'Table Grid'
        t.autofit = False
        widths = [Cm(1.6), Cm(2.4), Cm(4.6), Cm(5.4), Cm(3.8)]
        for tr in rows:
            cells = tr.find_all(['th', 'td'])
            rc = t.add_row().cells
            for i, c in enumerate(cells):
                rc[i].width = widths[i]
                p = rc[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                if c.name == 'th':
                    r = p.add_run(c.get_text(strip=True)); r.bold = True; r.font.size = Pt(9.5); shade(rc[i], 'E8EFE8')
                else:
                    render_inline(p, c, size=9.5)
        tbl = t._tbl
        tblPr = tbl.tblPr
        layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
        tblW = tblPr.find(qn('w:tblW'))
        if tblW is None: tblW = OxmlElement('w:tblW'); tblPr.append(tblW)
        tblW.set(qn('w:type'), 'dxa'); tblW.set(qn('w:w'), str(int(sum(w.twips for w in widths))))
        grid = tbl.tblGrid
        for gc in list(grid): grid.remove(gc)
        for w in widths:
            gc = OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(w.twips))); grid.append(gc)
        for i, w in enumerate(widths):
            for cell in t.columns[i].cells: cell.width = w
        doc.add_paragraph()
doc.save(out)
print('saved', out)
